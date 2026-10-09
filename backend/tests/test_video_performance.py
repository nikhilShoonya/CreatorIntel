"""Video Performance module: import, CRUD, daily refresh, history, discovery, sentiment, dashboard.

External HTTP is served by in-test fakes (httpx.MockTransport). Views change between
"days" so growth numbers can be verified exactly.
"""

import csv
import io
import json
import time
from datetime import datetime, timedelta, timezone

import httpx
import pandas as pd
import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import select

from app.config.settings import get_settings
from app.main import app
from app.models.db import session_scope
from app.services.llm_client import LLMClient
from app.video_performance.models import VtViewSnapshot
from app.video_performance.platforms import InstagramVideoClient, YouTubeVideoClient
from app.video_performance.sentiment import VideoSentimentAnalyzer
from app.video_performance.tracker import VideoTracker
from app.video_performance.urls import parse_instagram_handle, parse_video_link

YT_VIEWS = {"aaaaaaaaaaa": 2000, "bbbbbbbbbbb": 50_000}
NEW_UPLOADS: list[dict] = []  # videos "published" on the tracked channel after tracking started
IG_MEDIA = {
    "tradingtech31": [
        {"code": "Cabc12345", "views": 1000, "likes": 90, "comments": 10, "caption": "Bank Nifty levels for tomorrow"},
    ]
}


def now_iso() -> str:
    # a moment after "now" so it is clearly published after tracking started
    return (datetime.now(timezone.utc) + timedelta(seconds=2)).strftime("%Y-%m-%dT%H:%M:%SZ")


def youtube(request: httpx.Request) -> httpx.Response:
    path = request.url.path.rsplit("/", 1)[-1]
    params = request.url.params
    if path == "videos":
        items = []
        for vid in params["id"].split(","):
            if vid in YT_VIEWS:
                items.append({
                    "id": vid,
                    "snippet": {"title": f"Nifty Analysis {vid[:3]}", "description": "Market ka next move kya hoga? Detailed levels.",
                                "channelId": "UCaaaaaaaaaaaaaaaaaaaaaa", "channelTitle": "Nitin Nitro", "publishedAt": "2024-09-01T10:00:00Z"},
                    "statistics": {"viewCount": str(YT_VIEWS[vid]), "likeCount": str(YT_VIEWS[vid] // 20), "commentCount": "10"},
                })
        return httpx.Response(200, json={"items": items})
    if path == "channels":
        return httpx.Response(200, json={"items": [{
            "id": "UCaaaaaaaaaaaaaaaaaaaaaa", "snippet": {"title": "Nitin Nitro"},
            "contentDetails": {"relatedPlaylists": {"uploads": "UUaaaaaaaaaaaaaaaaaaaaaa"}},
        }]})
    if path == "playlistItems":
        return httpx.Response(200, json={"items": [
            {"contentDetails": {"videoId": v["id"], "videoPublishedAt": v["published"]}} for v in NEW_UPLOADS
        ] + [{"contentDetails": {"videoId": "bbbbbbbbbbb", "videoPublishedAt": "2023-01-01T00:00:00Z"}}]})
    return httpx.Response(404, json={"error": {"code": 404}})


def instagram(request: httpx.Request) -> httpx.Response:
    fields = request.url.params.get("fields", "")
    username = fields.split("username(", 1)[1].split(")", 1)[0]
    if username not in IG_MEDIA:
        return httpx.Response(400, json={"error": {"code": 110, "error_subcode": 2207013, "message": "Invalid user id"}})
    media = [{
        "id": m["code"], "permalink": f"https://www.instagram.com/reel/{m['code']}/", "media_type": "VIDEO",
        "media_product_type": "REELS", "timestamp": m.get("timestamp", "2024-09-01T10:00:00+0000"),
        "like_count": m["likes"], "comments_count": m["comments"], "view_count": m["views"], "caption": m["caption"],
    } for m in IG_MEDIA[username]]
    return httpx.Response(200, json={"business_discovery": {
        "id": "1784", "username": username, "name": "Trading Tech", "followers_count": 1000, "media": {"data": media},
    }})


def llm(request: httpx.Request) -> httpx.Response:
    body = json.loads(request.content)
    assert body["response_format"]["json_schema"]["name"] == "video_sentiment"
    return httpx.Response(200, json={"choices": [{"message": {"content": json.dumps({"sentiment": "Positive", "confidence": 0.9})}}]})


def client_for(handler) -> httpx.AsyncClient:
    return httpx.AsyncClient(transport=httpx.MockTransport(handler))


@pytest.fixture(scope="module")
def api():
    with TestClient(app) as test_client:
        settings = get_settings()
        app.state.video_tracker = VideoTracker(
            settings,
            youtube=YouTubeVideoClient(settings, client=client_for(youtube)),
            instagram=InstagramVideoClient(settings, client=client_for(instagram)),
            sentiment=VideoSentimentAnalyzer(settings, LLMClient(settings, client=client_for(llm))),
        )
        yield test_client


def wait_idle(api, timeout=10):
    deadline = time.time() + timeout
    while time.time() < deadline:
        items = api.get("/api/video-performance/videos", params={"page_size": 100}).json()["items"]
        if all(v["status"] not in ("Pending", "Processing") for v in items):
            return {v["video_identifier"]: v for v in items}
        time.sleep(0.05)
    raise AssertionError("videos still processing")


def backdate_snapshots(days: int = 1) -> None:
    """Pretend every existing snapshot was taken `days` earlier (simulates the next day)."""
    with session_scope() as db:
        for snapshot in db.scalars(select(VtViewSnapshot)):
            snapshot.captured_at = snapshot.captured_at - timedelta(days=days)


def test_url_parsing():
    assert parse_video_link("https://youtu.be/aaaaaaaaaaa?si=x").identifier == "aaaaaaaaaaa"
    assert parse_video_link("https://www.youtube.com/shorts/aaaaaaaaaaa").platform == "youtube"
    reel = parse_video_link("https://www.instagram.com/tradingtech31/reel/Cabc12345/?igsh=1")
    assert (reel.identifier, reel.owner_username) == ("Cabc12345", "tradingtech31")
    assert parse_video_link("https://www.instagram.com/tradingtech31/").ok is False
    assert parse_video_link("https://www.youtube.com/@nitinnitro").ok is False
    assert parse_instagram_handle("@TradingTech31") == "tradingtech31"
    assert parse_instagram_handle("Nitin Nitro") is None


def test_full_video_tracking_flow(api):
    frame = pd.DataFrame(
        [
            ["Nitin Nitro", "YouTube", "https://www.youtube.com/watch?v=aaaaaaaaaaa", ""],
            ["Nitin Nitro", "", "https://youtu.be/aaaaaaaaaaa", ""],  # duplicate in file
            ["Trading Tech", "Instagram", "https://www.instagram.com/reel/Cabc12345/", "@tradingtech31"],
            ["Some Person", "Instagram", "https://www.instagram.com/reel/Cnoowner1/", ""],  # owner unknown
            ["Broken", "", "not a link", ""],
        ],
        columns=["Creator Name", "Platform", "Video Link", "Username"],
    )
    buffer = io.BytesIO()
    frame.to_excel(buffer, index=False)
    res = api.post("/api/video-performance/uploads", files={"file": ("videos.xlsx", buffer.getvalue(), "application/octet-stream")})
    assert res.status_code == 201, res.text
    body = res.json()
    assert body["upload"]["added"] == 3 and body["upload"]["duplicates"] == 1 and body["upload"]["invalid"] == 1
    assert [r["status"] for r in body["rows"]] == ["added", "duplicate", "added", "added", "invalid"]

    videos = wait_idle(api)
    yt, ig, unknown = videos["aaaaaaaaaaa"], videos["Cabc12345"], videos["Cnoowner1"]
    assert yt["status"] == "Tracking" and yt["current_views"] == 2000 and yt["previous_views"] is None
    assert yt["engagement_rate"] == pytest.approx((100 + 10) / 2000 * 100, abs=0.01)
    assert yt["sentiment"] == "Positive"
    assert ig["status"] == "Tracking" and ig["current_views"] == 1000 and ig["engagement_rate"] == 10.0
    assert unknown["status"] == "Unsupported" and "username" in unknown["status_reason"]

    # The import-ready exports have the same four columns in both formats.
    csv_export = api.get("/api/video-performance/exports/csv", params={"layout": "upload"})
    excel_export = api.get("/api/video-performance/exports/excel", params={"layout": "upload"})
    assert csv_export.status_code == excel_export.status_code == 200
    csv_rows = list(csv.reader(io.StringIO(csv_export.content.decode("utf-8-sig"))))
    workbook = load_workbook(io.BytesIO(excel_export.content))
    sheet = workbook.active
    excel_rows = [[value if value is not None else "" for value in row] for row in sheet.values]
    assert csv_rows[0] == excel_rows[0] == ["Video Link", "Creator Name", "Platform", "Username"]
    assert sorted(csv_rows[1:]) == sorted(excel_rows[1:])
    assert any(row[0].endswith("/Cabc12345/") and row[3] == "tradingtech31" for row in csv_rows[1:])
    assert all(row[3] == "" for row in csv_rows[1:] if row[2] == "YouTube")
    # Shared export design: coloured header with filters, frozen header row, clickable links.
    assert workbook.sheetnames == ["Video list", "About"]  # data first, so it stays re-uploadable
    assert sheet["A1"].fill.fgColor.rgb == "004338CA" and sheet["A1"].font.bold
    assert sheet.freeze_panes == "A2"
    assert sheet.max_column == 4
    assert sheet.column_dimensions["A"].width >= 24
    assert sheet.auto_filter.ref == f"A1:D{sheet.max_row}"
    assert sheet["A2"].hyperlink is not None
    assert sheet["A2"].hyperlink.target == sheet["A2"].value
    ig_row = next(row for row in sheet.iter_rows(min_row=2) if row[2].value == "Instagram")
    assert ig_row[2].font.color.rgb == "00C13584"  # platform in its brand colour

    report = api.get("/api/video-performance/exports/excel")
    report_csv = api.get("/api/video-performance/exports/csv")
    assert report.status_code == report_csv.status_code == 200
    report_book = load_workbook(io.BytesIO(report.content))
    report_sheet = report_book.active
    report_csv_rows = list(csv.reader(io.StringIO(report_csv.content.decode("utf-8-sig"))))
    report_excel_rows = [[value if value is not None else "" for value in row] for row in report_sheet.values]
    assert report_csv_rows[0] == report_excel_rows[0]
    assert len(report_csv_rows) == len(report_excel_rows) == len(csv_rows)
    assert sorted((row[1], row[13]) for row in report_csv_rows[1:]) == sorted(
        (row[1], row[13]) for row in report_excel_rows[1:]
    )
    assert report_sheet.max_column == 18
    assert report_sheet["E1"].value == "Current Views"
    assert report_sheet["A1"].fill.fgColor.rgb == "004338CA"
    assert report_sheet.freeze_panes == "B2"  # header row + video title stay visible
    assert report_sheet.auto_filter.ref == f"A1:R{report_sheet.max_row}"
    assert report_sheet["P1"].value.startswith("Last Checked (") and report_sheet.column_dimensions["P"].width >= 19
    assert report_sheet["B2"].hyperlink.target == report_sheet["B2"].value
    tracked_row = next(r for r in report_sheet.iter_rows(min_row=2) if r[13].value == "Tracking")
    assert tracked_row[4].number_format == "#,##0"  # thousands separators on views
    assert tracked_row[10].number_format == '0.00"%"'  # engagement rate
    assert tracked_row[13].fill.fgColor.rgb == "00ECFDF5"  # status coloured like the app
    assert hasattr(tracked_row[15].value, "year")  # a real date (sortable), not text
    assert "Rows" in [c.value for c in report_book["About"]["A"]]
    assert all(value != "" for row in report_csv_rows[1:] for value in row)
    assert all(value != "" for row in report_excel_rows[1:] for value in row)
    unsupported_csv = next(row for row in report_csv_rows[1:] if "/Cnoowner1/" in row[1])
    unsupported_excel = next(row for row in report_excel_rows[1:] if "/Cnoowner1/" in row[1])
    for row in (unsupported_csv, unsupported_excel):
        assert row[4] == "N/A"  # current views
        assert row[5:8] == ["-", "-", "-"]  # previous views and growth
        assert row[11] == "N/A"  # sentiment

    # re-upload of the same video is not duplicated
    again = api.post(
        "/api/video-performance/uploads",
        files={"file": ("again.csv", b"Video Link\nhttps://www.youtube.com/watch?v=aaaaaaaaaaa\n", "text/csv")},
    ).json()
    assert again["rows"][0]["status"] == "already_tracked"

    # --- next day: daily refresh job computes gained / growth from yesterday's snapshot
    backdate_snapshots(1)
    YT_VIEWS["aaaaaaaaaaa"] = 2500
    IG_MEDIA["tradingtech31"][0]["views"] = 1500
    assert api.post("/api/video-performance/jobs/metrics_refresh/run").status_code == 202
    time.sleep(0.3)
    videos = wait_idle(api)
    yt = videos["aaaaaaaaaaa"]
    assert (yt["current_views"], yt["previous_views"], yt["views_gained"], yt["growth_pct"]) == (2500, 2000, 500, 25.0)
    assert videos["Cabc12345"]["views_gained"] == 500

    history = api.get(f"/api/video-performance/history/{yt['id']}").json()
    assert [s["views"] for s in history["snapshots"]] == [2500, 2000]
    assert history["snapshots"][0]["views_change"] == 500

    jobs = api.get("/api/video-performance/jobs").json()
    refresh = next(j for j in jobs if j["job_type"] == "metrics_refresh")
    assert refresh["last_status"] == "completed"

    # --- pause / resume / edit / retry
    assert api.post(f"/api/video-performance/videos/{yt['id']}/pause").status_code == 200
    assert api.get(f"/api/video-performance/videos/{yt['id']}").json()["status"] == "Paused"
    assert api.post(f"/api/video-performance/videos/{yt['id']}/resume").status_code == 202
    assert wait_idle(api)["aaaaaaaaaaa"]["status"] == "Tracking"

    fixed = api.put(f"/api/video-performance/videos/{unknown['id']}", json={"instagram_username": "tradingtech31"})
    assert fixed.status_code == 200
    assert wait_idle(api)["Cnoowner1"]["status_reason"] == "Reel not found on @tradingtech31's account"

    renamed = api.put(f"/api/video-performance/videos/{ig['id']}", json={"creator_name": "Trading Tech Official"})
    assert renamed.json()["creator_name"] == "Trading Tech Official"

    # --- manual add, duplicate manual add, invalid manual add
    added = api.post("/api/video-performance/videos", json={"video_url": "https://youtu.be/bbbbbbbbbbb", "creator_name": "Nitin Nitro"})
    assert added.status_code == 201
    assert api.post("/api/video-performance/videos", json={"video_url": "https://youtu.be/bbbbbbbbbbb"}).status_code == 409
    assert api.post("/api/video-performance/videos", json={"video_url": "https://twitter.com/x/status/1"}).status_code == 422

    # --- creator tracking + discovery of a newly published video
    creator = api.post("/api/video-performance/creators", json={"channel_url": "https://www.youtube.com/@nitinnitro", "creator_name": "Nitin Nitro"})
    assert creator.status_code == 201, creator.text
    creator_id = creator.json()["id"]
    time.sleep(0.3)
    listed = {c["id"]: c for c in api.get("/api/video-performance/creators").json()}
    assert listed[creator_id]["status"] == "Active"
    NEW_UPLOADS.append({"id": "ccccccccccc", "published": now_iso()})
    YT_VIEWS["ccccccccccc"] = 300
    assert api.post("/api/video-performance/jobs/creator_discovery/run").status_code == 202
    time.sleep(0.3)
    videos = wait_idle(api)
    assert videos["ccccccccccc"]["source"] == "discovered" and videos["ccccccccccc"]["current_views"] == 300
    assert "bbbbbbbbbbb" in videos and videos["bbbbbbbbbbb"]["source"] == "manual"  # old upload not re-added

    bad_creator = api.post("/api/video-performance/creators", json={"channel_url": "https://www.instagram.com/personal_acct/"})
    time.sleep(0.3)
    status = {c["id"]: c for c in api.get("/api/video-performance/creators").json()}[bad_creator.json()["id"]]
    assert status["status"] == "Unsupported" and "Professional" in status["status_reason"]

    # --- dashboard KPIs
    dash = api.get("/api/video-performance/dashboard").json()
    assert dash["total_videos"] == 5
    assert dash["youtube_videos"] == 3 and dash["instagram_videos"] == 2
    assert dash["new_videos_7d"] == 1
    assert dash["sentiment_overall"]["positive"] >= 3
    assert dash["top_performing"][0]["video_identifier"] == "bbbbbbbbbbb"
    assert dash["latest_detected"][0]["video_identifier"] == "ccccccccccc"

    # --- retry failed, bulk + delete
    assert api.post("/api/video-performance/videos/retry-failed").status_code == 202
    bulk = api.post("/api/video-performance/videos/bulk", json={"ids": [videos["ccccccccccc"]["id"]], "action": "pause"})
    assert bulk.json()["affected"] == 1
    assert api.delete(f"/api/video-performance/videos/{unknown['id']}").status_code == 200
    assert api.get(f"/api/video-performance/history/{unknown['id']}").status_code == 404
    assert api.delete(f"/api/video-performance/creators/{creator_id}").status_code == 200

    # --- isolation: Creator Analytics tables untouched
    assert api.get("/api/creators", params={"q": "ccccccccccc"}).json()["total"] == 0
